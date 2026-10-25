import datetime
import json
from types import SimpleNamespace

import openpyxl
import pytest

from app.api.documents import process_document_for_rag
from app.core.config import settings
from app.core.domain_profile import domain_profile
from app.core.limits import ADMIN_UPLOAD_EXTRACTION_LIMITS, MAX_RECORD_LABEL_CHARS
from app.rag.research_session import ResearchSession
from app.rag.tools import ResearchToolRegistry
from app.services.document_processor import DocumentProcessor
from app.services.structured_data import count_rendered_records, merge_record_parts, record_chunks
from app.services.structured_data.extract import extract_tree_text
from app.services.structured_data.models import WARNING_OVER_LIMIT, WARNING_PARSE_FAILED
from app.services.structured_data.render import RECORD_SEPARATOR, render_record

LABELS = ["title", "name", "author", "id"]
IGNORED = ["svg", "embedCode"]


@pytest.fixture(autouse=True)
def record_profile(monkeypatch):
    monkeypatch.setattr(domain_profile, "record_label_fields", LABELS)
    monkeypatch.setattr(domain_profile, "record_ignore_fields", IGNORED)


def tree(ext, text, limit=None):
    return extract_tree_text(ext, text, limit, LABELS, IGNORED)


def blocks(result):
    return result.text.split(RECORD_SEPARATOR)


def extract_file(tmp_path, name, content):
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")
    return DocumentProcessor.extract_document(str(path), "", ADMIN_UPLOAD_EXTRACTION_LIMITS)


def test_json_array_keeps_nested_fields_and_original_keys():
    data = [
        {
            "author": "alice",
            "content": "第一則",
            "postLink": "https://example.com/p/1",
            "tags": ["review", "good"],
            "likes": 12,
            "svg": "<svg><path/></svg>",
            "contact": {"email": "a@example.com"},
            "empty": None,
            "blank": "  ",
        },
        {"title": "退款流程", "items": [{"name": "a"}, {"name": "b"}], "flag": True, "ratio": 2.0},
        {"nested": {"only": "nested object"}},
        "純文字元素",
        None,
    ]

    result = tree(".json", json.dumps(data, ensure_ascii=False))

    assert (result.kind, result.structured_format, result.warning) == ("json", True, None)
    assert blocks(result) == [
        "[記錄 1]alice\nauthor: alice\ncontent: 第一則\npostLink: https://example.com/p/1\n"
        "tags: review, good\nlikes: 12\ncontact.email: a@example.com",
        "[記錄 2]退款流程\ntitle: 退款流程\nitems[0].name: a\nitems[1].name: b\nflag: true\nratio: 2",
        "[記錄 3]\nnested.only: nested object",
        "[記錄 4]\n純文字元素",
    ]
    assert [record.record_index for record in result.records] == [1, 2, 3, 4]
    assert result.records[0].links == ("https://example.com/p/1",)
    assert result.field_names[:3] == ("author", "content", "postLink")


def test_object_with_record_arrays_becomes_groups_and_one_context_record():
    data = {
        "total": 2,
        "source": "export",
        "data": [{"id": 7, "text": "甲"}, {"id": 8, "text": "乙"}],
        "users": [{"name": "喵"}],
        "tags": ["x", "y"],
    }

    result = tree(".json", json.dumps(data, ensure_ascii=False))

    assert blocks(result) == [
        "[上下文]\ntotal: 2\nsource: export\ntags: x, y",
        "[data 記錄 1]7\nid: 7\ntext: 甲",
        "[data 記錄 2]8\nid: 8\ntext: 乙",
        "[users 記錄 1]喵\nname: 喵",
    ]
    assert [record.group for record in result.records] == [None, "data", "data", "users"]
    assert [record.record_index for record in result.records] == [1, 2, 3, 4]


def test_object_without_record_arrays_is_a_single_record():
    data = {"site_name": "範例站", "max_upload_mb": 10, "maintenance": False, "contact": {"email": "s@example.com"}}

    result = tree(".json", json.dumps(data, ensure_ascii=False))

    assert blocks(result) == [
        "[記錄 1]\nsite_name: 範例站\nmax_upload_mb: 10\nmaintenance: false\ncontact.email: s@example.com"
    ]


def test_root_scalar_json_is_plain_text_without_warning():
    result = tree(".json", '"hello"')

    assert (result.kind, result.records, result.structured_format, result.warning) == ("plain", (), True, None)
    assert result.text == '"hello"'


def test_invalid_json_and_over_limit_fall_back_to_plain_text_with_warning():
    invalid = tree(".json", '[{"a": 1},]')
    over_limit = tree(".json", '[{"a": 1}]', limit=5)

    assert (invalid.kind, invalid.warning, invalid.text) == ("plain", WARNING_PARSE_FAILED, '[{"a": 1},]')
    assert (over_limit.kind, over_limit.warning, over_limit.text) == ("plain", WARNING_OVER_LIMIT, '[{"a": 1}]')
    assert invalid.structured_format and over_limit.structured_format


def test_json_with_byte_order_mark_is_parsed():
    result = tree(".json", '﻿[{"title": "a"}]')

    assert blocks(result) == ["[記錄 1]a\ntitle: a"]


def test_jsonl_lines_are_records_and_one_bad_line_falls_back():
    good = tree(".jsonl", '{"title": "a"}\n\n{"title": "b", "n": 1}\n')
    bad = tree(".jsonl", '{"title": "a"}\n{bad\n')

    assert good.kind == "jsonl"
    assert blocks(good) == ["[記錄 1]a\ntitle: a", "[記錄 2]b\ntitle: b\nn: 1"]
    assert (bad.kind, bad.warning) == ("plain", WARNING_PARSE_FAILED)


def test_js_declaration_needs_strict_json_and_never_warns():
    declared = tree(".js", 'const posts = [{"title": "x"}];')
    literal = tree(".js", "const posts = [{title: 'x'}];")
    code = tree(".ts", "function f() { return 1; }")

    assert declared.kind == "json" and blocks(declared) == ["[記錄 1]x\ntitle: x"]
    assert (literal.kind, literal.warning, literal.structured_format) == ("plain", None, False)
    assert (code.kind, code.warning, code.structured_format) == ("plain", None, False)


def test_yaml_is_parsed_with_dates_in_iso_format():
    text = "owner: hr\nitems:\n  - title: 年會\n    date: 2026-08-01\n    when: 2026-08-01 09:30:00\n  - title: 尾牙\n"

    result = tree(".yaml", text)

    assert result.kind == "yaml"
    assert blocks(result) == [
        "[上下文]\nowner: hr",
        "[items 記錄 1]年會\ntitle: 年會\ndate: 2026-08-01\nwhen: 2026-08-01 09:30:00",
        "[items 記錄 2]尾牙\ntitle: 尾牙",
    ]


def test_yaml_alias_expansion_and_syntax_errors_fall_back():
    levels = ["a: &a [x, x, x, x, x, x, x, x, x]"]
    for name, previous in zip("bcde", "abcd"):
        levels.append(f"{name}: &{name} [" + ", ".join([f"*{previous}"] * 9) + "]")
    bomb = tree(".yaml", "\n".join(levels) + "\n")
    broken = tree(".yml", "a: [1, 2\n")

    assert (bomb.kind, bomb.warning) == ("plain", WARNING_PARSE_FAILED)
    assert (broken.kind, broken.warning) == ("plain", WARNING_PARSE_FAILED)


def test_noise_values_are_replaced_and_ignored_fields_are_dropped():
    data = [{
        "title": "x",
        "logo": "data:image/png;base64," + "A" * 120,
        "html": "前<svg><path/></svg>後",
        "embedCode": "<blockquote>embed</blockquote>",
    }]

    result = tree(".json", json.dumps(data, ensure_ascii=False))

    assert blocks(result) == ["[記錄 1]x\ntitle: x\nlogo: [BASE64_IMAGE_OMITTED]\nhtml: 前[SVG_ICON_OMITTED]後"]


def test_label_is_truncated_in_the_header_only():
    title = "長" * (MAX_RECORD_LABEL_CHARS + 10)

    record = tree(".json", json.dumps([{"title": title}], ensure_ascii=False)).records[0]

    assert record.label == "長" * MAX_RECORD_LABEL_CHARS + "…"
    assert record.fields == (("title", title),)


def test_csv_rows_become_records_with_spreadsheet_row_numbers(tmp_path):
    content = "﻿title,,date,title\n\n退款,x,2026-08-01,重複\n,,,\n請假,,2026-08-02,,額外\n"

    result = extract_file(tmp_path, "rows.csv", content)

    assert result.kind == "csv"
    assert blocks(result) == [
        "[第 3 列]退款\ntitle: 退款\n2: x\ndate: 2026-08-01\ntitle_2: 重複",
        "[第 5 列]請假\ntitle: 請假\ndate: 2026-08-02\n5: 額外",
    ]
    assert [record.row for record in result.records] == [3, 5]


def test_csv_with_only_a_header_stays_plain_text(tmp_path):
    result = extract_file(tmp_path, "header.csv", "title,date\n")

    assert (result.kind, result.structured_format, result.warning) == ("plain", True, None)
    assert result.text == "title | date"


def test_xlsx_rows_use_sheet_names_column_letters_and_iso_dates(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "員工"
    for column, value in enumerate(["name", None, "到職日", "amount"], start=1):
        sheet.cell(row=2, column=column, value=value)
    for column, value in enumerate(["小明", "x", datetime.datetime(2026, 8, 1), 3.0], start=1):
        sheet.cell(row=3, column=column, value=value)
    for column, value in enumerate(["小華", None, datetime.datetime(2026, 8, 2, 9, 30), 2.5], start=1):
        sheet.cell(row=5, column=column, value=value)
    workbook.create_sheet("空白").cell(row=1, column=1, value="only header")
    path = tmp_path / "staff.xlsx"
    workbook.save(path)

    result = DocumentProcessor.extract_document(str(path), "", ADMIN_UPLOAD_EXTRACTION_LIMITS)

    assert result.kind == "xlsx"
    assert blocks(result) == [
        "[員工 第 3 列]小明\nname: 小明\nB: x\n到職日: 2026-08-01\namount: 3",
        "[員工 第 5 列]小華\nname: 小華\n到職日: 2026-08-02 09:30:00\namount: 2.5",
    ]
    assert [(record.sheet, record.row, record.record_index) for record in result.records] == [("員工", 3, 1), ("員工", 5, 2)]


def test_long_record_is_split_into_parts_with_a_repeated_header():
    data = [{"title": "長記錄", "a": "甲" * 50, "b": "乙" * 50, "c": "丙" * 300, "url": "https://example.com/a"}]
    record = tree(".json", json.dumps(data, ensure_ascii=False)).records[0]

    chunks = record_chunks(record, 120)

    total = len(chunks)
    assert total > 2
    for index, (text, metadata) in enumerate(chunks, start=1):
        assert len(text) <= 120
        assert text.startswith(f"[記錄 1]長記錄（{index}/{total}）\n")
        assert metadata == {
            "record_index": 1,
            "part": index,
            "part_count": total,
            "preserve_whole": True,
            "label": "長記錄",
            "links": ["https://example.com/a"],
        }
    merged = merge_record_parts([text for text, _ in chunks])
    assert merged.split("\n", 1)[0] == "[記錄 1]長記錄"
    assert "a: " + "甲" * 50 in merged
    assert "丙" * 300 in merged.replace("\n", "")


def test_long_value_pieces_keep_sentence_punctuation_at_the_end():
    body = "報名期間為八月。" + "行程包含步道與海岸，住宿為雙人房。" * 20
    record = tree(".json", json.dumps([{"title": "旅遊", "content": body}], ensure_ascii=False)).records[0]

    bodies = [text.split("\n", 1)[1] for text, _ in record_chunks(record, 120)]

    assert len(bodies) > 2
    assert not any(piece.startswith("。") for piece in bodies)
    assert "".join(bodies).replace("title: 旅遊", "").replace("content: ", "") == body


def test_short_record_is_one_chunk_without_part_suffix():
    record = tree(".json", '[{"title": "短"}]').records[0]

    assert record_chunks(record, 800) == [(render_record(record), {
        "record_index": 1, "part": 1, "part_count": 1, "preserve_whole": True, "label": "短",
    })]


def test_faq_detection_is_skipped_for_structured_formats():
    rag_system = SimpleNamespace(chunk_size=800)
    base = {"source": "doc", "document_id": 1}
    faq_text = "Q: 特休幾天？\nA: 依年資。"

    data_docs, data_qa, data_records = process_document_for_rag(
        DocumentProcessor.stored_text_extraction("faq.json", faq_text), base, rag_system
    )
    text_docs, text_qa, text_records = process_document_for_rag(
        DocumentProcessor.stored_text_extraction("faq.txt", faq_text), base, rag_system
    )
    record_docs, record_qa, record_count = process_document_for_rag(
        tree(".json", '[{"title": "a"}, {"title": "b"}]'), base, rag_system
    )

    assert (len(data_docs), data_qa, data_records) == (1, 0, 0)
    assert (len(text_docs), text_qa, text_records) == (1, 1, 0)
    assert (len(record_docs), record_qa, record_count) == (2, 0, 2)
    assert record_docs[1].metadata == {
        "source": "doc", "document_id": 1, "record_index": 2, "part": 1, "part_count": 1,
        "preserve_whole": True, "label": "b",
    }


def indexed_chunks(*records_with_ids):
    documents = []
    chunk_id = 0
    for document_id, record, chunk_size in records_with_ids:
        for text, metadata in record_chunks(record, chunk_size):
            chunk_id += 1
            documents.append(SimpleNamespace(page_content=text, metadata={
                **metadata, "chunk_id": chunk_id, "document_id": document_id, "source": f"doc{document_id}.json",
            }))
    return documents


@pytest.mark.anyio
async def test_filter_counts_split_records_once_and_matches_across_parts():
    data = [
        {"title": "公告", "timestamp": "2026-08-01", "body": "甲" * 200 + "關鍵字", "url": "https://example.com/a"},
        {"title": "其他", "timestamp": "2026-08-02", "body": "沒有"},
    ]
    records = tree(".json", json.dumps(data, ensure_ascii=False)).records
    documents = indexed_chunks((1, records[0], 100), (1, records[1], 100))
    documents.append(SimpleNamespace(page_content="一般文字 2026-08-03 關鍵字", metadata={"chunk_id": 99}))
    store = SimpleNamespace(documents=documents)
    registry = ResearchToolRegistry(retriever=SimpleNamespace(vector_store=store))

    result = await registry.filter_and_count_records(date_range="2026-08", keyword="關鍵字")

    assert len([doc for doc in documents if doc.metadata.get("record_index") == 1]) > 2
    assert result["total_count"] == 2
    record, plain = result["records"]
    assert (record["chunk_id"], record["record_index"], record["label"]) == (1, 1, "公告")
    assert record["links"] == ["https://example.com/a"]
    assert record["content"].startswith("timestamp: 2026-08-01\nbody:")
    assert "公告" not in record["content"] and "https://example.com/a" not in record["content"]
    assert (plain["chunk_id"], plain["label"], plain["links"]) == (99, None, [])


def test_record_links_from_filter_results_allow_web_fetch(monkeypatch):
    monkeypatch.setattr(settings, "BLOCK_WEB_TOOLS_AFTER_KB", False)
    session = ResearchSession(["統計貼文"])
    output = {"records": [{"chunk_id": 1, "source": "a.json", "content": "內容", "links": ["https://example.com/post/1"]}]}

    session.record_tool_output("filter_and_count_records", output)

    assert session.refuse_tool_call("web_fetch", {"url": "https://example.com/post/1"}) is None


def test_fallback_summary_uses_the_extraction_or_counts_rendered_headers():
    extraction = tree(".json", '[{"title": "a", "n": 1}, {"title": "b"}]')

    with_extraction = DocumentProcessor._generate_fallback_summary("a.json", extraction.text, "", extraction)
    text_only = DocumentProcessor._generate_fallback_summary("a.json", extraction.text)

    assert with_extraction == "收錄 2 筆 JSON 結構化記錄，欄位包括：title、n。"
    assert text_only == "收錄約 2 筆結構化記錄。"
    assert count_rendered_records("[第 3 列]a\nx: 1\n\n---\n\n[工作表 1 第 4 列]\ny: 2\n\n---\n\n[上下文]\nz: 3") == 3
    assert count_rendered_records("[重點]\n內容") == 0
