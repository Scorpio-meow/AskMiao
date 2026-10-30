import dataclasses
import importlib.util
import io
import json
import time
from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument
from sqlalchemy import event
from starlette.datastructures import Headers, UploadFile

from app.api import admin as admin_api
from app.api import documents as documents_api
from app.core.config import settings
from app.core.domain_profile import domain_profile
from app.models import Document as DbDocument, RagChunk
from app.rag import tools as tools_module
from app.rag.indices import bm25_store as bm25_store_module
from app.services import document_removal
from app.services.pdf_service import PDFService
from test_rag_index_consistency import (
    LEAVE,
    OVERTIME,
    add_indexed_document,
    assert_indices_match_database,
    bm25_sources,
    index_dir,
    make_rag,
    rag,
    session_factory,
)

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def upload_dir(tmp_path, monkeypatch):
    directory = tmp_path / "uploads"
    directory.mkdir()
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(directory))
    monkeypatch.setattr(documents_api, "UPLOAD_DIR", str(directory))
    return directory


@pytest.fixture
def served_rag(rag, monkeypatch):
    monkeypatch.setattr(documents_api, "get_rag_system", lambda: rag)
    monkeypatch.setattr(document_removal, "get_rag_system", lambda: rag)
    return rag


@pytest.fixture
def offline_summary(monkeypatch):
    async def fake_summary(filename, content, file_type="", extraction=None):
        return f"summary of {filename}"

    monkeypatch.setattr(documents_api.DocumentProcessor, "generate_document_summary_async", fake_summary)


def text_upload(filename, text):
    return file_upload(filename, text.encode("utf-8"), "text/plain")


def file_upload(filename, data, content_type):
    return UploadFile(
        file=io.BytesIO(data),
        filename=filename,
        headers=Headers({"content-type": content_type}),
    )


def scanned_pdf(width):
    doc = pymupdf.open()
    doc.new_page(width=width, height=842)
    data = doc.tobytes()
    doc.close()
    return data


def protected_pdf():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "confidential overtime policy")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="user-secret", owner_pw="owner-secret")
    doc.close()
    return data


def docx_bytes(text):
    document = DocxDocument()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def upload(session_factory, *files):
    with session_factory() as db:
        response = await documents_api.upload_document(
            file=list(files), db=db, user_id=1, current_user={"username": "admin"}
        )
    return response["results"]


def stored_documents(session_factory):
    with session_factory() as session:
        return [(doc.id, doc.filename, doc.is_processed) for doc in session.query(DbDocument).order_by(DbDocument.id)]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "delete_endpoint",
    [documents_api.delete_document, admin_api.delete_document],
    ids=["documents", "admin"],
)
async def test_delete_endpoints_remove_upload_file(delete_endpoint, served_rag, session_factory, upload_dir):
    (upload_dir / "overtime.txt").write_text(OVERTIME, encoding="utf-8")
    (upload_dir / "leave.txt").write_text(LEAVE, encoding="utf-8")
    add_indexed_document(served_rag, session_factory, 1, "overtime.txt", OVERTIME)
    add_indexed_document(served_rag, session_factory, 2, "leave.txt", LEAVE)

    with session_factory() as db:
        await delete_endpoint(1, db=db, current_user={"username": "admin"})

    assert [path.name for path in upload_dir.iterdir()] == ["leave.txt"]
    assert [doc_id for doc_id, _, _ in stored_documents(session_factory)] == [2]
    assert bm25_sources(served_rag, "approval") == []
    assert_indices_match_database(served_rag, session_factory)


@pytest.mark.asyncio
async def test_delete_endpoint_fails_when_index_removal_fails(served_rag, make_rag, session_factory, upload_dir, monkeypatch):
    (upload_dir / "overtime.txt").write_text(OVERTIME, encoding="utf-8")
    add_indexed_document(served_rag, session_factory, 1, "overtime.txt", OVERTIME)

    def locked_writer(index):
        raise OSError("bm25 index is locked")

    with monkeypatch.context() as patched:
        patched.setattr(bm25_store_module, "AsyncWriter", locked_writer)
        patched.setattr(bm25_store_module.time, "sleep", lambda seconds: None)
        with session_factory() as db:
            with pytest.raises(OSError, match="bm25 index is locked"):
                await documents_api.delete_document(1, db=db, current_user={"username": "admin"})

    assert stored_documents(session_factory) == []
    assert list(upload_dir.iterdir()) == []
    served_rag.bm25_store.bm25_searcher.close()
    assert_indices_match_database(make_rag(), session_factory)


@pytest.mark.asyncio
async def test_startup_indexes_documents_left_unprocessed_by_a_crashed_upload(
    served_rag, session_factory, upload_dir, monkeypatch
):
    monkeypatch.setattr(domain_profile, "record_label_fields", ["title"])
    posts = [{"title": "停車規定", "body": "parking permit"}, {"title": "門禁規定", "body": "badge access"}]
    (upload_dir / "posts.json").write_text(json.dumps(posts, ensure_ascii=False), encoding="utf-8")
    with session_factory() as session:
        session.add_all([
            DbDocument(id=1, filename="overtime.txt", content=OVERTIME, file_type="text/plain", is_processed=False),
            DbDocument(id=2, filename="leave.txt", content=LEAVE, file_type="text/plain", is_processed=True),
            DbDocument(id=3, filename="posts.json", content="舊格式內容", file_type="application/json", is_processed=False),
        ])
        session.commit()

    with session_factory() as db:
        await documents_api.index_unprocessed_documents(db)

    assert stored_documents(session_factory) == [(1, "overtime.txt", True), (2, "leave.txt", True), (3, "posts.json", True)]
    assert bm25_sources(served_rag, "approval") == ["overtime.txt"]
    assert bm25_sources(served_rag, "seniority") == []
    records = sorted(
        (doc.metadata["record_index"], doc.metadata["label"]) for doc in served_rag.documents if doc.metadata["document_id"] == 3
    )
    assert records == [(1, "停車規定"), (2, "門禁規定")]
    assert_indices_match_database(served_rag, session_factory)

@pytest.mark.asyncio
async def test_upload_failure_after_commit_removes_row_chunks_and_upload_file(
    served_rag, session_factory, upload_dir, offline_summary, monkeypatch
):
    index_documents = served_rag.add_documents

    def index_then_fail(docs):
        index_documents(docs)
        raise RuntimeError("index save failed")

    monkeypatch.setattr(served_rag, "add_documents", index_then_fail)

    with session_factory() as db:
        response = await documents_api.upload_document(
            file=[text_upload("overtime.txt", OVERTIME)], db=db, user_id=1, current_user={"username": "admin"}
        )

    assert [(item["status"], item["http_status"]) for item in response["results"]] == [("failed", 500)]
    assert list(upload_dir.iterdir()) == []
    assert stored_documents(session_factory) == []
    with session_factory() as session:
        assert session.query(RagChunk).count() == 0
    assert bm25_sources(served_rag, "approval") == []
    assert_indices_match_database(served_rag, session_factory)


@pytest.mark.asyncio
async def test_upload_failure_keeps_upload_file_of_earlier_document(
    served_rag, session_factory, upload_dir, offline_summary, monkeypatch
):
    validate_file_type = documents_api.DocumentProcessor.validate_file_type

    def validate_or_crash(content_type, filename=None):
        if filename == "broken.txt":
            raise RuntimeError("validator crashed")
        return validate_file_type(content_type, filename)

    monkeypatch.setattr(documents_api.DocumentProcessor, "validate_file_type", validate_or_crash)

    with session_factory() as db:
        response = await documents_api.upload_document(
            file=[text_upload("overtime.txt", OVERTIME), text_upload("broken.txt", LEAVE)],
            db=db, user_id=1, current_user={"username": "admin"},
        )

    assert [(item["status"], item["http_status"]) for item in response["results"]] == [("success", 201), ("failed", 500)]
    assert [path.name for path in upload_dir.iterdir()] == ["overtime.txt"]
    assert [(filename, is_processed) for _, filename, is_processed in stored_documents(session_factory)] == [
        ("overtime.txt", True)
    ]
    assert_indices_match_database(served_rag, session_factory)


@pytest.mark.asyncio
async def test_upload_does_not_take_the_filename_of_a_document_whose_file_is_missing(
    served_rag, session_factory, upload_dir, offline_summary
):
    add_indexed_document(served_rag, session_factory, 1, "overtime.txt", OVERTIME)

    results = await upload(session_factory, text_upload("overtime.txt", LEAVE))

    assert [(item["status"], item["filename"]) for item in results] == [("success", "overtime_1.txt")]
    assert [path.name for path in upload_dir.iterdir()] == ["overtime_1.txt"]
    assert [(doc_id, filename) for doc_id, filename, _ in stored_documents(session_factory)] == [
        (1, "overtime.txt"), (2, "overtime_1.txt")
    ]


@pytest.mark.asyncio
async def test_renamed_upload_stays_within_the_filename_byte_limit(
    served_rag, session_factory, upload_dir, offline_summary
):
    from app.core.limits import MAX_FILENAME_BYTES

    name = "ab" + "假" * ((MAX_FILENAME_BYTES - len("ab.txt")) // 3) + ".txt"
    add_indexed_document(served_rag, session_factory, 1, name, OVERTIME)

    results = await upload(session_factory, text_upload(name, LEAVE))

    assert results[0]["status"] == "success"
    renamed = results[0]["filename"]
    assert renamed.endswith("_1.txt") and len(renamed.encode("utf-8")) <= MAX_FILENAME_BYTES
    assert (upload_dir / renamed).exists()


@pytest.mark.asyncio
async def test_failed_commit_does_not_break_the_rest_of_the_batch(
    served_rag, session_factory, upload_dir, offline_summary
):
    def fail_first_insert(mapper, connection, target):
        if target.filename == "overtime.txt":
            raise RuntimeError("database is locked")

    event.listen(DbDocument, "before_insert", fail_first_insert)
    try:
        results = await upload(session_factory, text_upload("overtime.txt", OVERTIME), text_upload("leave.txt", LEAVE))
    finally:
        event.remove(DbDocument, "before_insert", fail_first_insert)

    assert [(item["status"], item["http_status"]) for item in results] == [("failed", 500), ("success", 201)]
    assert [path.name for path in upload_dir.iterdir()] == ["leave.txt"]
    assert [(filename, is_processed) for _, filename, is_processed in stored_documents(session_factory)] == [
        ("leave.txt", True)
    ]
    assert_indices_match_database(served_rag, session_factory)


@pytest.mark.asyncio
@pytest.mark.parametrize("vision_available", [False, True], ids=["ocr-not-configured", "ocr-failed"])
async def test_scanned_pdfs_without_any_text_are_rejected_instead_of_indexed(
    vision_available, served_rag, session_factory, upload_dir, offline_summary, monkeypatch
):
    monkeypatch.setattr(PDFService, "_is_vision_available", classmethod(lambda cls: vision_available))
    monkeypatch.setattr(
        PDFService, "_ocr_page_sync", classmethod(lambda cls, img, page_num: f"[第 {page_num} 頁為掃描圖片或複雜版面，已記錄]")
    )

    results = await upload(
        session_factory,
        file_upload("scan_a.pdf", scanned_pdf(595), "application/pdf"),
        file_upload("scan_b.pdf", scanned_pdf(612), "application/pdf"),
    )

    assert [(item["status"], item["http_status"]) for item in results] == [("failed", 400), ("failed", 400)]
    assert all("沒有可提取的文字" in item["detail"] for item in results)
    assert list(upload_dir.iterdir()) == []
    assert stored_documents(session_factory) == []


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "make_upload, detail",
    [
        (lambda: file_upload("protected.pdf", protected_pdf(), "application/pdf"), "PDF 文件被密碼保護，無法讀取"),
        (lambda: file_upload("fake.pdf", b"plain text pretending to be a PDF", "application/pdf"), "安全驗證失敗: 上傳的檔案未通過安全檢查"),
        (lambda: file_upload("broken.docx", b"PK\x03\x04 not a zip archive", DOCX_TYPE), "無效的 Office 文件格式"),
        (lambda: file_upload("handbook.docx", docx_bytes(OVERTIME), DOCX_TYPE), "文件內含的檔案數超過上限"),
        (lambda: text_upload("長" * 90 + ".txt", OVERTIME), "檔名驗證失敗: 無效的檔名或不允許的副檔名"),
    ],
    ids=["password-protected", "forged-header", "invalid-office-file", "over-limit", "filename-too-long-in-bytes"],
)
async def test_rejected_uploads_report_why(
    make_upload, detail, served_rag, session_factory, upload_dir, offline_summary, monkeypatch
):
    monkeypatch.setattr(
        documents_api,
        "ADMIN_UPLOAD_EXTRACTION_LIMITS",
        dataclasses.replace(documents_api.ADMIN_UPLOAD_EXTRACTION_LIMITS, max_ooxml_members=1),
    )

    results = await upload(session_factory, make_upload())

    assert [(item["status"], item["detail"], item["http_status"]) for item in results] == [("failed", detail, 400)]
    assert list(upload_dir.iterdir()) == []
    assert stored_documents(session_factory) == []


def load_reprocess_script():
    path = Path(__file__).resolve().parent.parent / "scripts" / "reprocess_existing_docs.py"
    spec = importlib.util.spec_from_file_location("reprocess_existing_docs_cli", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.asyncio
async def test_reprocess_script_invalidates_knowledge_base_description(
    served_rag, session_factory, upload_dir, offline_summary, monkeypatch
):
    add_indexed_document(served_rag, session_factory, 1, "overtime.txt", OVERTIME)
    add_indexed_document(served_rag, session_factory, 2, "leave.txt", LEAVE)
    with session_factory() as session:
        session.get(DbDocument, 1).description = "管理員修訂過的摘要"
        session.commit()
    script = load_reprocess_script()
    monkeypatch.setattr(script, "SessionLocal", session_factory)
    monkeypatch.setattr(tools_module, "_kb_description_cache", (time.monotonic(), "stale description"))

    await script.reprocess_all()

    assert tools_module._kb_description_cache is None
    with session_factory() as session:
        assert [doc.description for doc in session.query(DbDocument).order_by(DbDocument.id)] == [
            "管理員修訂過的摘要", "summary of leave.txt"
        ]
    assert bm25_sources(served_rag, "approval") == ["overtime.txt"]
    assert_indices_match_database(served_rag, session_factory)
