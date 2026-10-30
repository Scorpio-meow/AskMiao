#!/usr/bin/env python3
import os
import csv
import logging
import hashlib
import re
import zipfile
from typing import Optional
from lxml import etree
from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from pptx import Presentation
from pptx.shapes.group import GroupShape
import openpyxl
from app.core.config import settings
from app.core.domain_profile import domain_profile
from app.core.html_text import extract_text_and_title
from app.core.limits import MAX_OOXML_COMPRESSION_RATIO, MAX_SUMMARY_FIELD_NAMES, OOXML_RATIO_CHECK_MIN_BYTES, ExtractionLimits
from app.services.structured_data import (
    KIND_LABELS,
    TREE_TEXT_EXTENSIONS,
    ExtractionResult,
    clean_noise,
    count_rendered_records,
    extract_csv,
    extract_tree_text,
    extract_xlsx,
    is_data_format,
    plain_result,
)
logger = logging.getLogger(__name__)
# 摘要啟發式只看每行開頭的這些字元
MAX_TOC_LINE_CHARS = 300
OOXML_CONTENT_TYPES_PART = "[Content_Types].xml"
# 與 python-docx／python-pptx 相同：不展開實體、不連網
OOXML_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)


class DocumentRejected(ValueError):
    pass


class DocumentLimitExceeded(DocumentRejected):
    """文件超過本次抽取的資源預算；訊息只描述哪一項上限，可以原樣告知使用者"""


def _is_xml_content_type(content_type: str) -> bool:
    lowered = content_type.lower()
    return lowered.endswith("+xml") or lowered.endswith("/xml")


def _ooxml_xml_bytes(archive: zipfile.ZipFile, members: list, max_xml_bytes: int) -> int:
    """依 [Content_Types].xml 加總會被當成 XML 解析的成員大小（另含 .rels 與該檔本身）。

    python-docx／python-pptx 依內容類型而非副檔名決定是否把部件建成 DOM，因此改名不能繞過這個預算"""
    try:
        content_types_info = archive.getinfo(OOXML_CONTENT_TYPES_PART)
    except KeyError as e:
        raise DocumentRejected("無效的 Office 文件格式") from e
    if content_types_info.file_size > max_xml_bytes:
        raise DocumentLimitExceeded("文件的 XML 內容超過解析上限")
    try:
        root = etree.fromstring(archive.read(content_types_info), parser=OOXML_XML_PARSER)
    except etree.XMLSyntaxError as e:
        raise DocumentRejected("無效的 Office 文件格式") from e
    defaults = {}
    overrides = {}
    for element in root:
        if not isinstance(element.tag, str):
            continue
        name = etree.QName(element).localname
        content_type = element.get("ContentType")
        extension = element.get("Extension")
        part_name = element.get("PartName")
        if name == "Default" and content_type is not None and extension is not None:
            defaults[extension.lower()] = content_type
        elif name == "Override" and content_type is not None and part_name is not None:
            overrides[part_name.lower()] = content_type
    total = content_types_info.file_size
    for info in members:
        if info.filename == OOXML_CONTENT_TYPES_PART:
            continue
        name = info.filename.lower()
        content_type = overrides.get("/" + name)
        if content_type is None:
            content_type = defaults.get(os.path.splitext(name)[1].lstrip("."))
        if name.endswith(".rels") or (content_type is not None and _is_xml_content_type(content_type)):
            total += info.file_size
    return total


CODE_OR_DATA_EXTENSIONS = (".js", ".ts", ".jsx", ".tsx", ".json", ".jsonl", ".py", ".yaml", ".yml", ".ini", ".env", ".sql")


class DocumentProcessor:
    
    FILE_SIGNATURES = {
        'application/pdf': [b'%PDF'],
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document': [b'PK\x03\x04'],
        'application/vnd.openxmlformats-officedocument.presentationml.presentation': [b'PK\x03\x04'],
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': [b'PK\x03\x04'],
        'text/plain': [],
        'text/markdown': [],
        'text/x-markdown': [],
        'text/csv': [],
        'application/json': [],
        'text/json': [],
        'text/html': [],
        'text/xml': [],
        'application/xml': [],
        'text/yaml': [],
        'application/x-yaml': []
    }
    SUPPORTED_EXTENSIONS = {
        '.txt', '.md', '.markdown', '.pdf', '.docx', '.pptx', '.xlsx',
        '.csv', '.json', '.jsonl', '.yaml', '.yml', '.xml', '.html', '.htm', '.log',
        '.py', '.js', '.ts', '.tsx', '.jsx', '.java', '.cpp', '.c', '.sql',
        '.sh', '.ini', '.env'
    }
    
    @staticmethod
    def validate_file_header(file_path: str, content_type: str) -> bool:
        signatures = DocumentProcessor.FILE_SIGNATURES.get(content_type)
        if signatures is None or not signatures:
            return True
        
        try:
            with open(file_path, 'rb') as f:
                header = f.read(512)
                
            for signature in signatures:
                if header.startswith(signature):
                    return True
            
            logger.warning(f"檔案頭部驗證失敗: {file_path}, 聲稱類型: {content_type}")
            raise ValueError(
                f"檔案頭部與聲稱的類型不符。"
                f"這可能是一個偽造的檔案或損壞的檔案。"
            )
            
        except Exception as e:
            logger.error(f"驗證檔案頭部時發生錯誤: {e}")
            raise
    
    @staticmethod
    def calculate_file_hash(file_path: str) -> str:
        sha256_hash = hashlib.sha256()
        with open(file_path, "rb") as f:
            for byte_block in iter(lambda: f.read(4096), b""):
                sha256_hash.update(byte_block)
        return sha256_hash.hexdigest()
    
    @staticmethod
    def extract_text_from_file(file_path: str, content_type: str, limits: ExtractionLimits) -> Optional[str]:
        """limits：本次抽取的資源預算，管理員上傳與聊天附件各用 app.core.limits 中的一組"""
        return DocumentProcessor.extract_document(file_path, content_type, limits).text

    @staticmethod
    def stored_text_extraction(filename: str, content: str) -> ExtractionResult:
        return plain_result(content, is_data_format(filename), None)

    @staticmethod
    def extract_document(file_path: str, content_type: str, limits: ExtractionLimits) -> ExtractionResult:
        try:
            ext = os.path.splitext(file_path)[1].lower()
            
            if ext == ".docx":
                content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif ext == ".pptx":
                content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            elif ext == ".xlsx":
                content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            elif ext == ".pdf":
                content_type = "application/pdf"
            elif ext in ('.md', '.markdown'):
                content_type = "text/markdown"
            elif ext == '.csv':
                content_type = "text/csv"
            elif ext == '.json':
                content_type = "application/json"
            elif ext in ('.html', '.htm'):
                content_type = "text/html"
            DocumentProcessor.validate_file_header(file_path, content_type)
            file_size = os.path.getsize(file_path)
            max_size = int(settings.MAX_FILE_SIZE_MB) * 1024 * 1024
            if file_size > max_size:
                raise DocumentRejected(f"檔案大小 ({file_size} bytes) 超過限制")
            if ext == ".pdf" or content_type == "application/pdf":
                return plain_result(DocumentProcessor._extract_from_pdf(file_path, limits.max_ocr_pages), False, None)
            elif ext == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                DocumentProcessor._check_ooxml_archive(file_path, limits, builds_dom=True)
                return plain_result(DocumentProcessor._extract_from_docx(file_path), False, None)
            elif ext == ".pptx" or content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
                DocumentProcessor._check_ooxml_archive(file_path, limits, builds_dom=True)
                return plain_result(DocumentProcessor._extract_from_pptx(file_path), False, None)
            elif ext == ".xlsx" or content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                # openpyxl 以 read_only 串流讀取工作表，不會整份建成 DOM
                DocumentProcessor._check_ooxml_archive(file_path, limits, builds_dom=False)
                return extract_xlsx(
                    file_path, domain_profile.record_label_fields, lambda: DocumentProcessor._extract_from_xlsx(file_path)
                )
            elif ext == ".csv" or content_type == "text/csv":
                return extract_csv(
                    file_path,
                    domain_profile.record_label_fields,
                    lambda: DocumentProcessor._extract_from_csv(file_path),
                    lambda: DocumentProcessor._extract_from_txt(file_path),
                )
            elif ext in TREE_TEXT_EXTENSIONS:
                return extract_tree_text(
                    ext,
                    DocumentProcessor._extract_from_txt(file_path),
                    limits.max_structured_parse_chars,
                    domain_profile.record_label_fields,
                    domain_profile.record_ignore_fields,
                )
            elif ext in CODE_OR_DATA_EXTENSIONS:
                return plain_result(clean_noise(DocumentProcessor._extract_from_txt(file_path)), False, None)
            elif ext in (".html", ".htm") or content_type == "text/html":
                return plain_result(DocumentProcessor._extract_from_html(file_path), False, None)
            else:
                return plain_result(DocumentProcessor._extract_from_txt(file_path), False, None)
        except Exception as e:
            logger.error(f"提取文本時發生錯誤: {e}")
            raise
    
    @staticmethod
    def _check_ooxml_archive(file_path: str, limits: ExtractionLimits, builds_dom: bool) -> None:
        """OOXML 是 ZIP：解析前先檢查成員數、宣告的解壓大小與壓縮比，擋下壓縮炸彈。

        zipfile 讀取成員時不會超過宣告的大小，因此宣告值可以當作上界。壓縮比除了逐一成員，
        也看整個檔案，拆成許多小成員的炸彈同樣會被擋下。builds_dom 為 True 的格式另外限制
        會被建成 lxml DOM 的 XML 總量：DOM 的記憶體是 XML 大小的十幾到幾十倍"""
        try:
            with zipfile.ZipFile(file_path) as archive:
                members = archive.infolist()
                if len(members) > limits.max_ooxml_members:
                    raise DocumentLimitExceeded("文件內含的檔案數超過上限")
                total = 0
                total_compressed = 0
                for info in members:
                    total += info.file_size
                    total_compressed += info.compress_size
                    if total > limits.max_ooxml_uncompressed_bytes:
                        raise DocumentLimitExceeded("文件解壓後的大小超過上限")
                    if info.file_size > OOXML_RATIO_CHECK_MIN_BYTES and info.file_size > info.compress_size * MAX_OOXML_COMPRESSION_RATIO:
                        raise DocumentLimitExceeded("文件的壓縮比異常，疑似壓縮炸彈")
                if total > OOXML_RATIO_CHECK_MIN_BYTES and total > total_compressed * MAX_OOXML_COMPRESSION_RATIO:
                    raise DocumentLimitExceeded("文件的壓縮比異常，疑似壓縮炸彈")
                if builds_dom and _ooxml_xml_bytes(archive, members, limits.max_ooxml_xml_bytes) > limits.max_ooxml_xml_bytes:
                    raise DocumentLimitExceeded("文件的 XML 內容超過解析上限")
        except zipfile.BadZipFile as e:
            raise DocumentRejected("無效的 Office 文件格式") from e
    @staticmethod
    def _extract_from_txt(file_path: str) -> str:
        encodings = ['utf-8', 'utf-8-sig', 'big5', 'gbk', 'gb2312', 'latin1']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                if content.strip():
                    logger.info(f"成功使用 {encoding} 編碼讀取文字文件")
                    return content
            except UnicodeDecodeError:
                continue
        
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        return content
    @staticmethod
    def _extract_from_csv(file_path: str) -> str:
        encodings = ['utf-8', 'utf-8-sig', 'big5', 'gbk', 'gb2312', 'latin1']
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding, newline='') as f:
                    reader = csv.reader(f)
                    rows = []
                    for row in reader:
                        if any(cell.strip() for cell in row):
                            rows.append(" | ".join(row))
                    if rows:
                        return "\n".join(rows)
            except Exception:
                continue
        return DocumentProcessor._extract_from_txt(file_path)
    @staticmethod
    def _extract_from_html(file_path: str) -> str:
        raw_text = DocumentProcessor._extract_from_txt(file_path)
        # 線性掃描：標準函式庫 html.parser 遇到大量未閉合的標籤會呈二次時間
        _, _, blocks = extract_text_and_title(raw_text)
        extracted = "\n".join(blocks)
        return extracted or raw_text
    
    @staticmethod
    def _extract_from_pdf(file_path: str, max_ocr_pages: Optional[int]) -> str:
        from app.services.pdf_service import PDFService
        return PDFService.extract_text_robust(file_path, max_ocr_pages=max_ocr_pages)
    
    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        doc = DocxDocument(file_path)
        text_content = []
        
        for block in doc.iter_inner_content():
            if isinstance(block, DocxTable):
                for row in block.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_content.append(" | ".join(row_text))
            elif block.text.strip():
                text_content.append(block.text)
        
        if not text_content:
            raise DocumentRejected("DOCX 文件中沒有可提取的文本內容")
        
        full_text = "\n".join(text_content)
        logger.info(f"成功從 DOCX 提取 {len(doc.paragraphs)} 個段落和 {len(doc.tables)} 個表格，共 {len(full_text)} 字符")
        return full_text
    @staticmethod
    def _extract_from_pptx(file_path: str) -> str:
        prs = Presentation(file_path)
        slides_text = []
        for idx, slide in enumerate(prs.slides):
            slide_lines = [f"--- 投影片 {idx + 1} ---"]
            slide_lines.extend(DocumentProcessor._pptx_shape_lines(slide.shapes))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_lines.append(f"備忘稿：{notes}")
            if len(slide_lines) > 1:
                slides_text.append("\n".join(slide_lines))
        
        if not slides_text:
            raise DocumentRejected("PPTX 文件中沒有可提取的文本內容")
        
        full_text = "\n\n".join(slides_text)
        logger.info(f"成功從 PPTX 提取 {len(prs.slides)} 頁投影片，共 {len(full_text)} 字符")
        return full_text
    @staticmethod
    def _pptx_shape_lines(shapes) -> list:
        lines = []
        for shape in shapes:
            if isinstance(shape, GroupShape):
                lines.extend(DocumentProcessor._pptx_shape_lines(shape.shapes))
            elif shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        lines.append(" | ".join(row_cells))
        return lines
    @staticmethod
    def _extract_from_xlsx(file_path: str) -> str:
        # read_only 以串流方式讀取工作表，不把整份活頁簿展開到記憶體
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        sheets_text = []
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            sheet_lines = [f"--- 工作表: {sheetname} ---"]
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() != "" for cell in row):
                    row_str = [str(c).strip() if c is not None else "" for c in row]
                    sheet_lines.append(" | ".join(row_str))
            if len(sheet_lines) > 1:
                sheets_text.append("\n".join(sheet_lines))
        
        if not sheets_text:
            raise DocumentRejected("XLSX 文件中沒有可提取的文本內容")
        
        full_text = "\n\n".join(sheets_text)
        logger.info(f"成功從 XLSX 提取 {len(wb.sheetnames)} 個工作表，共 {len(full_text)} 字符")
        wb.close()
        return full_text
    
    @staticmethod
    def validate_file_type(content_type: str, file_path: str = None) -> bool:
        if content_type in DocumentProcessor.FILE_SIGNATURES:
            return True
        if file_path:
            ext = os.path.splitext(file_path)[1].lower()
            if ext in DocumentProcessor.SUPPORTED_EXTENSIONS:
                return True
        return False
    @staticmethod
    def _records_overview(extraction: Optional[ExtractionResult]) -> Optional[str]:
        if extraction is None or not extraction.records:
            return None
        overview = f"{len(extraction.records)} 筆 {KIND_LABELS[extraction.kind]} 結構化記錄"
        fields = extraction.field_names[:MAX_SUMMARY_FIELD_NAMES]
        if fields:
            overview += f"，欄位包括：{'、'.join(fields)}"
        return overview
    @classmethod
    async def generate_document_summary_async(
        cls, filename: str, content: str, file_type: str = "", extraction: Optional[ExtractionResult] = None
    ) -> str:
        """
        使用 LLM 智能生成專業、結構化、流暢且具高資訊密度的繁體中文「文件主題與大綱摘要」。
        若 LLM 呼叫超時或發生異常，自動降級調用高品質的啟發式規則分析器。
        """
        if not content or not content.strip():
            return f"收錄內部文件《{filename}》。"
        overview = cls._records_overview(extraction)
        stripped = content.strip()
        if len(stripped) <= 2800:
            sample_content = stripped
        else:
            first_part = stripped[:1500]
            mid_start = max(0, (len(stripped) // 2) - 400)
            mid_part = stripped[mid_start:mid_start + 800]
            last_part = stripped[-500:]
            sample_content = f"{first_part}\n\n[...中略...]\n\n{mid_part}\n\n[...下略...]\n\n{last_part}"
        prompt = (
            "你是一個頂級的知識庫文件分析與大綱生成專家。請根據以下文件的檔名與內文摘錄，"
            "產出一份高水準、專業、客觀且條理分明的「文件主題與大綱摘要」（繁體中文，約 70~150 字）。\n\n"
            "[嚴格準則]\n"
            "1. 核心定位：準確指明文件類型與核心主題（例如：產品規格說明書、宣傳型錄DM、LINE工作討論紀錄、社群貼文資料集、技術手冊等）。\n"
            "2. 內容大綱：提煉出主要涵蓋的章節、核心架構、功能特色或討論焦點。\n"
            "3. 噪音過濾：嚴格禁止將通訊對話人名/任務待辦（如「1. @某某」）、純頁碼、版權宣告或公司電話地址誤當作章節標題。\n"
            "4. 輸出要求：請直接輸出摘要文字本身，語句通順精煉，嚴禁客套話或前言標籤。\n\n"
            f"[文件檔名]{filename}\n"
            f"[文件類型]{file_type}\n"
            + (f"[結構化資料]{overview}\n" if overview else "")
            + f"[內容摘錄]\n{sample_content}"
        )
        try:
            from app.core.llm_client import call_llm, resolve_model_name, system_endpoint
            from app.api.tags import models_and_default
            _, default_model = await models_and_default()
            model = resolve_model_name(default_model or None)
            response = await call_llm(
                messages=[{"role": "user", "content": prompt}],
                endpoint=system_endpoint(model),
                model_name=model
            )
            cleaned = response.strip()
            for prefix in ("好的，", "這是該文件的摘要：", "這是一份", "文件摘要：", "摘要：", "以下是"):
                if cleaned.startswith(prefix) and len(cleaned) > len(prefix) + 10:
                    cleaned = cleaned[len(prefix):].strip()
            if cleaned and len(cleaned) >= 20:
                return cleaned
        except Exception as e:
            logger.warning(f"使用 LLM 生成文件大綱摘要失敗 ({e})，切換至規則解析器降級處理")
        return cls._generate_fallback_summary(filename, content, file_type, extraction)
    @classmethod
    def generate_document_summary(cls, filename: str, content: str, file_type: str = "") -> str:
        """同步版只用規則摘要：呼叫端持有知識庫描述的快取鎖，等模型回應會讓等這把鎖的聊天請求一起卡住"""
        return cls._generate_fallback_summary(filename, content, file_type)
    @classmethod
    def _generate_fallback_summary(
        cls, filename: str, content: str, file_type: str = "", extraction: Optional[ExtractionResult] = None
    ) -> str:
        """高品質啟發式規則分析器"""
        if not content or not content.strip():
            return f"收錄內部文件《{filename}》。"
        raw_lines = [line.strip() for line in content.strip().splitlines() if line.strip()]
        if not raw_lines:
            return f"收錄內部文件《{filename}》。"
        overview = cls._records_overview(extraction)
        if overview:
            return f"收錄 {overview}。"
        record_count = count_rendered_records(content)
        if record_count > 0:
            return f"收錄約 {record_count} 筆結構化記錄。"
        header_preview = content[:2500]
        is_line_file = filename.startswith('[LINE]') or '對話' in filename or '通訊' in filename
        line_date_matches = len(re.findall(r'\d{4}[./-]\d{2}[./-]\d{2}\s+(?:星期|週|Mon|Tue|Wed|Thu|Fri|Sat|Sun)', header_preview))
        chat_time_user_matches = len(re.findall(r'\n\d{1,2}:\d{2}\s+[\u4e00-\u9fa5a-zA-Z0-9_]{2,12}', header_preview))
        
        summary_rules = domain_profile.summary_fallback
        if is_line_file or (line_date_matches >= 1 and chat_time_user_matches >= 2):
            members = set(re.findall(r'\n\d{1,2}:\d{2}\s+([\u4e00-\u9fa5a-zA-Z0-9_]{2,10})', header_preview))
            noise_words = set(summary_rules.chat_member_noise_words)
            members_filtered = [m for m in members if m not in noise_words and not m.isdigit()][:5]
            members_str = '、'.join(members_filtered) if members_filtered else '專案團隊成員'
            
            topics = []
            for kw in summary_rules.chat_topic_keywords:
                if kw in content and kw not in topics:
                    topics.append(kw)
            topic_str = '、'.join(topics[:5]) if topics else '專案工作事項'
            return f"本文件為內部通訊工作討論紀錄（參與人員包括：{members_str}），主要聚焦討論 {topic_str} 等相關任務與進度追蹤。"
        cleaned_lines = []
        noise_prefixes = ("---", "[", ">", "http://", "https://", "www.", "tel:", "fax:", "email:", "@")
        boilerplate_prefixes = tuple(p.lower() for p in summary_rules.boilerplate_line_prefixes)
        for line in raw_lines:
            line_str = line.strip()
            if any(line_str.startswith(np) for np in noise_prefixes):
                continue
            if line_str.lower().startswith(boilerplate_prefixes):
                continue
            if line_str in ("[本頁為圖檔或無可提取純文字]", "[SVG_ICON_OMITTED]", "[BASE64_IMAGE_OMITTED]"):
                continue
            cleaned_lines.append(line_str)
        if not cleaned_lines:
            return f"收錄內部文件《{filename}》。"
        def strip_toc_leader(text: str) -> str:
            """移除目錄行尾的「.....12」引導線與頁碼；以字元掃描取代含重疊量詞的正規式"""
            without_page = text.rstrip("0123456789")
            end = len(without_page)
            while end > 0 and (without_page[end - 1] in ".·_" or without_page[end - 1].isspace()):
                end -= 1
            if len(without_page) - end >= 3:
                return without_page[:end]
            return text
        def clean_toc_line(text: str) -> str:
            cleaned = text[:MAX_TOC_LINE_CHARS].lstrip("#*- •\t ")
            cleaned = strip_toc_leader(cleaned).strip()
            cleaned = re.sub(r'^(?:[0-9]+(?:\.[0-9]+)*\.?|[一二三四五六七八九十]+[、\.]|第[0-9一二三四五六七八九十]+[章節點條項篇]|[\[【][^\]】]+[\]】])\s*', '', cleaned).strip()
            return re.sub(r'\s+', ' ', cleaned)
        headers = []
        for line in cleaned_lines[:120]:
            cleaned_h = clean_toc_line(line)
            if len(cleaned_h) < 3 or len(cleaned_h) > 35:
                continue
            if re.search(r'^(?:第[0-9一二三四五六七八九十]+[章篇]|(?:[0-9]+|[一二三四五六七八九十]+)[、\.])\s*[^\d]', line):
                if cleaned_h not in headers:
                    headers.append(cleaned_h)
            elif line.startswith('#'):
                if cleaned_h not in headers:
                    headers.append(cleaned_h)
        title_cand = filename.rsplit('.', 1)[0]
        if headers:
            top_headers = headers[:5]
            headers_str = '、'.join(top_headers)
            return f"本文件為《{title_cand}》，核心涵蓋章節與重點包括：{headers_str} 等專業內容。"
        if "," in cleaned_lines[0] or "\t" in cleaned_lines[0] or "|" in cleaned_lines[0]:
            cols = [c.strip() for c in cleaned_lines[0].replace("|", ",").split(",") if c.strip()]
            if len(cols) >= 2:
                cols_str = "、".join(cols[:8])
                return f"結構化表格數據，收錄約 {len(cleaned_lines) - 1} 筆資料，包含欄位：{cols_str}。"
        snippet = " ".join(cleaned_lines[:4])[:120].replace("  ", " ")
        return f"本文件為《{title_cand}》，重點收錄：{snippet}...等內部專業資料。"
