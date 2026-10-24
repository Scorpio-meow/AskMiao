import csv
import itertools
from typing import Callable, Iterable, Iterator, List, Optional, Sequence, Set, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from .models import StructuredRecord
from .render import row_header
from .values import is_link, make_label, render_scalar, unique_links

TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "big5", "gbk", "gb2312", "latin1")
BOM = "﻿"


class _ColumnNames:
    def __init__(self, header: Sequence[str], fallback_name: Callable[[int], str]):
        self.fallback_name = fallback_name
        self.names: List[str] = []
        self.seen: Set[str] = set()
        for index, name in enumerate(header):
            self._add(name or fallback_name(index))

    def _add(self, name: str) -> None:
        candidate = name
        suffix = 2
        while candidate in self.seen:
            candidate = f"{name}_{suffix}"
            suffix += 1
        self.seen.add(candidate)
        self.names.append(candidate)

    def name(self, index: int) -> str:
        while index >= len(self.names):
            self._add(self.fallback_name(len(self.names)))
        return self.names[index]


def _csv_column_name(index: int) -> str:
    return str(index + 1)


def _xlsx_column_name(index: int) -> str:
    return get_column_letter(index + 1)


def records_from_rows(
    rows: Iterable[Tuple[int, Sequence[str]]],
    label_fields: Sequence[str],
    sheet: Optional[str],
    fallback_name: Callable[[int], str],
    counter: Iterator[int],
) -> List[StructuredRecord]:
    columns: Optional[_ColumnNames] = None
    records: List[StructuredRecord] = []
    for row_number, cells in rows:
        if not any(cells):
            continue
        if columns is None:
            columns = _ColumnNames(cells, fallback_name)
            continue
        fields = tuple((columns.name(index), cell) for index, cell in enumerate(cells) if cell)
        values = {name: value for name, value in fields}
        label = next((make_label(values[name]) for name in label_fields if values.get(name)), None)
        records.append(StructuredRecord(
            record_index=next(counter),
            header=row_header(row_number, sheet),
            label=label,
            fields=fields,
            links=unique_links(value for _, value in fields if is_link(value)),
            group=None,
            sheet=sheet,
            row=row_number,
        ))
    return records


def read_csv_rows(file_path: str) -> List[List[str]]:
    for encoding in TEXT_ENCODINGS:
        try:
            with open(file_path, "r", encoding=encoding, newline="") as f:
                return list(csv.reader(f))
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("csv", b"", 0, 1, "無法以支援的編碼讀取 CSV")


def _csv_cells(row: Sequence[str], first: bool) -> List[str]:
    cells = [render_scalar(cell) or "" for cell in row]
    if first and cells and cells[0].startswith(BOM):
        cells[0] = cells[0][len(BOM):].strip()
    return cells


def records_from_csv(file_path: str, label_fields: Sequence[str]) -> List[StructuredRecord]:
    rows = read_csv_rows(file_path)
    numbered = ((number, _csv_cells(row, number == 1)) for number, row in enumerate(rows, start=1))
    return records_from_rows(numbered, label_fields, None, _csv_column_name, itertools.count(1))


def _xlsx_rows(sheet) -> Iterator[Tuple[int, List[str]]]:
    for number, values in enumerate(sheet.iter_rows(min_row=1, values_only=True), start=1):
        yield number, [render_scalar(value) or "" for value in values]


def records_from_xlsx(file_path: str, label_fields: Sequence[str]) -> List[StructuredRecord]:
    workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        counter = itertools.count(1)
        records: List[StructuredRecord] = []
        for sheet_name in workbook.sheetnames:
            records.extend(records_from_rows(
                _xlsx_rows(workbook[sheet_name]), label_fields, sheet_name, _xlsx_column_name, counter
            ))
        return records
    finally:
        workbook.close()
